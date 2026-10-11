"""One-time extractor: workbook -> seed/ + test/fixtures.json. Usage: python3 tools/extract_seed.py book.xlsm"""
import json, re, sys
from openpyxl import load_workbook
SRC = sys.argv[1]; v = load_workbook(SRC, data_only=True)
norm = lambda s: re.sub(r'\s+', ' ', str(s if s is not None else '')).strip()
slug = lambda s: re.sub(r'[^a-z0-9]+', '-', s.lower()).strip('-') or 'x'
ws = v['Linked Move Types']; types = []
for r in range(4, 19):
    n = norm(ws.cell(r, 3).value); t = ws.cell(r, 4).value
    types.append({'id': slug(n), 'name': n, 'time': t if isinstance(t, (int, float)) else None})
ids = {t['id'] for t in types}
for t_ in types:
    if t_['id'] == 'other': t_['name'] = 'Manual'  # shown as Manual in the app
json.dump({'moveTypes': types}, open('seed/library.json', 'w'), indent=1)

ws = v['Motion Reference']; secs = []; cur = None
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 1).value
    if a is None: cur = None; continue
    if isinstance(ws.cell(r, 2).value, str) and str(a).isupper():
        mc = max(c for c in range(2, 18) if ws.cell(r, c).value is not None)
        cur = {'title': str(a).title(), 'headers': [norm(ws.cell(r, c).value) for c in range(2, mc + 1)], 'rows': []}; secs.append(cur)
    elif cur: cur['rows'].append([norm(a)] + [ws.cell(r, c).value for c in range(2, len(cur['headers']) + 2)])
json.dump(secs, open('seed/motion_reference.json', 'w'), indent=0)

fixtures = []
for name, sample in [('Fiber Splicing', 'sample-fiber-splicing'), ('Fiber Splicing (2)', 'sample-fiber-splicing-2'), ('Cell 10', None)]:
    ws = v[name]
    rc = next(r for r in range(8, 60) if ws.cell(r, 10).value == 'Cycletime')
    num2id = {int(ws.cell(r, 1).value): f's{r}' for r in range(8, rc) if isinstance(ws.cell(r, 1).value, (int, float))}
    steps, expect = [], []
    for r in range(8, rc):
        g = lambda c: ws.cell(r, c).value
        tid = slug(norm(g(3))) if norm(g(3)) else ''
        steps.append({'id': f's{r}', 'task': '' if name == 'Cell 10' else norm(g(2)), 'typeId': tid if tid in ids else '',
            'dist': g(4), 'vel': g(5), 'acc': g(6), 'settle': g(7), 'ovr': g(9),
            'deps': [num2id[int(g(c))] for c in (13, 14, 15) if isinstance(g(c), (int, float)) and int(g(c)) in num2id]})
        expect.append({'start': g(11), 'finish': g(12)})
    L = lambda k: ws.cell(rc + k, 12).value
    fixtures.append({'name': name, 'steps': steps, 'expect': expect, 'cycle': L(0), 'sf': L(1), 'nUp': L(3), 'uph': L(5)})
    if sample:
        keep = [s for s in steps if s['task']]
        json.dump({'id': sample, 'customer': '', 'job': name + ' (sample)', 'station': '', 'jobNo': '', 'originator': '', 'rev': '',
                   'sf': L(1), 'nUp': L(3), 'notes': '', 'steps': keep}, open(f'seed/projects/{sample}.json', 'w'), indent=1)
json.dump(fixtures, open('test/fixtures.json', 'w'))
print('ok', len(types), 'types', len(secs), 'ref sections', [(f['name'], f['cycle'], f['uph']) for f in fixtures])
